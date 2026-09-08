from __future__ import annotations

import math
import os
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.graphics import renderPDF, renderSVG
from reportlab.graphics.shapes import Drawing, Group, Line, Path as RPath, Polygon, Rect, String
from reportlab.lib.colors import HexColor, white


ROOT = Path(__file__).resolve().parents[1]
REFERENCE = ROOT / "Example" / "System Per Tab Explaination .docx"
OUTPUT = ROOT / "Documentation"
DOCX_OUT = OUTPUT / "Multimedia Management System Per Tab Explanation.docx"
ICON = ROOT / "public" / "web-app-manifest-512x512.png"


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def format_table(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row_idx, row in enumerate(table.rows):
        for idx, (cell, width) in enumerate(zip(row.cells, widths)):
            cell.width = width
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width.inches * 1440)))
            tc_w.set(qn("w:type"), "dxa")
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                for run in p.runs:
                    run.font.name = "Times New Roman"
                    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), "Times New Roman")
                    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), "Times New Roman")
                    run.font.size = Pt(8)
                    if row_idx == 0:
                        run.bold = True
        if row_idx == 0:
            set_repeat_table_header(row)
            for cell in row.cells:
                set_cell_shading(cell, "D9E2F3")


def clear_paragraph(paragraph):
    for child in list(paragraph._p):
        paragraph._p.remove(child)


def add_text(doc, text, style="Normal", bold_label=None):
    p = doc.add_paragraph(style=style)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if bold_label and text.startswith(bold_label):
        r = p.add_run(bold_label)
        r.bold = True
        p.add_run(text[len(bold_label):])
    else:
        p.add_run(text)
    for run in p.runs:
        run.font.name = "Times New Roman"
        run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), "Times New Roman")
        run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), "Times New Roman")
        run.font.color.rgb = RGBColor(0, 0, 0)
        if style == "Normal" or style.startswith("List"):
            run.font.size = Pt(10)
        elif style == "Heading 1":
            run.font.size = Pt(14)
            run.bold = True
        elif style == "Heading 2":
            run.font.size = Pt(11)
            run.bold = True
    return p


def add_bullets(doc, items):
    for item in items:
        add_text(doc, item, "List Bullet")


def new_numbering_id(doc):
    numbering = doc.part.numbering_part.element
    nums = numbering.findall(qn("w:num"))
    ids = [int(n.get(qn("w:numId"))) for n in nums]
    base = next((n for n in nums if n.get(qn("w:numId")) == "5"), nums[0])
    abstract = base.find(qn("w:abstractNumId")).get(qn("w:val"))
    new_id = max(ids) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(new_id))
    abstract_el = OxmlElement("w:abstractNumId")
    abstract_el.set(qn("w:val"), abstract)
    num.append(abstract_el)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:startOverride")
    start.set(qn("w:val"), "1")
    override.append(start)
    num.append(override)
    numbering.append(num)
    return new_id


def add_steps(doc, items):
    num_id = new_numbering_id(doc)
    for item in items:
        p = add_text(doc, item, "List Number")
        p_pr = p._p.get_or_add_pPr()
        num_pr = p_pr.get_or_add_numPr()
        ilvl = num_pr.get_or_add_ilvl()
        ilvl.set(qn("w:val"), "0")
        num = num_pr.get_or_add_numId()
        num.set(qn("w:val"), str(num_id))


def add_section(doc, number, title, purpose, info, workflows):
    add_text(doc, f"{number}. {title}", "Heading 1")
    add_text(doc, f"Purpose: {purpose}", bold_label="Purpose:")
    add_text(doc, "Information Displayed", "Heading 2")
    add_bullets(doc, info)
    for heading, steps in workflows:
        add_text(doc, heading, "Heading 2")
        add_steps(doc, steps)


def build_docx():
    OUTPUT.mkdir(parents=True, exist_ok=True)
    doc = Document(str(REFERENCE))

    for style_name in ("Normal", "Heading 1", "Heading 2", "List Bullet", "List Number", "TOC 1", "TOC 2", "TOC 3"):
        if style_name in doc.styles:
            style = doc.styles[style_name]
            style.font.name = "Times New Roman"
            style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), "Times New Roman")
            style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), "Times New Roman")
            style.font.color.rgb = RGBColor(0, 0, 0)
    doc.styles["Normal"].font.size = Pt(10)
    doc.styles["Heading 1"].font.size = Pt(14)
    doc.styles["Heading 1"].font.bold = True
    doc.styles["Heading 2"].font.size = Pt(11)
    doc.styles["Heading 2"].font.bold = True

    # Replace the retained cover image and title in place.
    cover_image = doc.paragraphs[6]
    clear_paragraph(cover_image)
    cover_image.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover_image.add_run().add_picture(str(ICON), width=Inches(2.15))

    title = doc.paragraphs[10]
    clear_paragraph(title)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Multimedia Management System")
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(24)

    doc.paragraphs[24].text = "Tab Explanation and Step-by-Step User Guide"
    doc.paragraphs[24].paragraph_format.page_break_before = True
    doc.paragraphs[24].runs[0].font.name = "Times New Roman"
    doc.paragraphs[24].runs[0].font.size = Pt(14)
    doc.paragraphs[25].text = (
        "This guide explains each tab and major workspace in the Multimedia Management System, "
        "what information appears there, who normally uses it, and the exact steps for common work. "
        "It is written for standard users, super administrators, and recipients of viewer or editor sharing links."
    )
    doc.paragraphs[26].text = (
        "Reference style: module purpose, information displayed, main actions, workflow, and controls. "
        "This version covers file organization, previews, versions, sharing, storage, recovery, sessions, and permissions."
    )

    anchor = doc.paragraphs[26]._p
    body = doc._element.body
    node = anchor.getnext()
    while node is not None and node.tag != qn("w:sectPr"):
        nxt = node.getnext()
        body.remove(node)
        node = nxt

    for footer in doc.sections[0].footer.paragraphs:
        if "Credit Monitoring System User Guide" in footer.text:
            footer.text = "Multimedia Management System User Guide"
            footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in footer.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(8)

    add_text(doc, "System Flow", "Heading 1")
    add_text(
        doc,
        "End-to-end process: Sign in -> Open My Drive or another tab -> Upload or select an item -> "
        "Organize, preview, version, or share it -> Background processing completes when required -> "
        "Use Recent, Starred, Shared, Storage, or Trash to find and manage the item later."
    )

    add_text(doc, "Sidebar Structure Overview", "Heading 1")
    add_text(doc, "This is the practical navigation map for the left sidebar and account menu. Use it to understand where common actions begin.")
    tree = [
        "Left Sidebar",
        "|-- New",
        "|   |-- New folder",
        "|   |-- File upload",
        "|   `-- Folder upload",
        "|-- My Drive",
        "|   |-- Root folders and nested folders",
        "|   |-- Grid or list layout",
        "|   `-- Rename, move, copy, download, share, star, details, trash",
        "|-- Recent",
        "|   `-- Recently opened files",
        "|-- Starred",
        "|   `-- Favorite files and folders",
        "|-- Shared",
        "|   `-- Items shared with the signed-in user",
        "|-- Users, super administrator only",
        "|   |-- Add account",
        "|   |-- Edit account and role",
        "|   `-- Delete account",
        "|-- Trash",
        "|   |-- Restore selected items",
        "|   `-- Delete permanently or apply a bulk action",
        "`-- Storage",
        "    |-- Usage totals and category breakdown",
        "    `-- Filter, sort, preview, or download files",
        "Profile Menu",
        "|-- My Drive",
        "|-- Profile",
        "|-- Active Devices",
        "|-- Users, super administrator only",
        "`-- Sign out",
    ]
    for line in tree:
        p = add_text(doc, line)
        p.paragraph_format.left_indent = Inches(0.15 if line.startswith("|") or line.startswith("`") or line.startswith(" ") else 0)
        for r in p.runs:
            r.font.size = Pt(8)
            if line in ("Left Sidebar", "Profile Menu"):
                r.bold = True

    add_text(doc, "Navigation And Access Summary", "Heading 2")
    nav_rows = [
        ("Area", "Main purpose", "Typical access"),
        ("My Drive", "Upload, create, browse, organize, preview, version, and share owned items.", "Signed-in users; super admins can see all owned content."),
        ("Recent", "Return to files recorded as recently opened.", "Signed-in users."),
        ("Starred", "Collect favorite files and folders for quick access.", "Signed-in users."),
        ("Shared", "Browse files and folders explicitly shared with the account.", "Signed-in users, according to viewer/editor permission."),
        ("Trash", "Restore retained items or permanently remove them.", "Item owner or super admin."),
        ("Storage", "Review storage totals, file categories, and large/recent files.", "Signed-in users; super admins see system-wide totals."),
        ("Users", "Create, edit, role-manage, and delete accounts.", "Super administrator only."),
        ("Preview", "Open supported media and document formats without downloading first.", "Authorized signed-in users or valid share recipients."),
        ("Profile / Devices", "Update the account and manage active sign-in sessions.", "Signed-in user."),
    ]
    table = doc.add_table(rows=1, cols=3, style="Table Grid")
    for idx, val in enumerate(nav_rows[0]):
        table.rows[0].cells[idx].text = val
    for row in nav_rows[1:]:
        cells = table.add_row().cells
        for idx, val in enumerate(row):
            cells[idx].text = val
    format_table(table, [Inches(1.1), Inches(3.25), Inches(2.15)])

    add_text(doc, "Role Access", "Heading 1")
    add_text(doc, "Access is controlled by the signed-in account role, item ownership, explicit share permission, and the settings on a public sharing link.")
    role_rows = [
        ("Role", "Can access", "Can change"),
        ("Super administrator", "All drive content, storage analytics, shared items, previews, profile, devices, and Users.", "Owned or accessible items plus user accounts and roles; cannot delete the currently signed-in account."),
        ("Standard user", "Own drive content and items shared with the account.", "Own items and shared items that grant editor access."),
        ("Viewer link", "The shared file or folder while the link is valid and any password/restriction check passes.", "No content changes; download only when the owner allows it."),
        ("Editor link", "The shared folder while the link is valid and any password/restriction check passes.", "Upload files and create folders inside the shared folder; existing owner controls remain protected."),
    ]
    table = doc.add_table(rows=1, cols=3, style="Table Grid")
    for idx, val in enumerate(role_rows[0]):
        table.rows[0].cells[idx].text = val
    for row in role_rows[1:]:
        cells = table.add_row().cells
        for idx, val in enumerate(row):
            cells[idx].text = val
    format_table(table, [Inches(1.25), Inches(2.55), Inches(2.7)])

    add_section(doc, 1, "My Drive",
        "My Drive is the main workspace for files and folders. It combines navigation, upload, organization, previews, sharing, versions, and item details.",
        [
            "Breadcrumbs for the current folder path and expandable root folders in the sidebar.",
            "Grid and list layouts for folders and files.",
            "File name, type, size, owner or sharing state, modified date, and processing state where applicable.",
            "New folder, file upload, and folder upload actions.",
            "Context actions for rename, move, copy, download, star, share, details, versions, and trash.",
            "Upload progress and background-finalization status for large or staged uploads.",
        ],
        [
            ("How To Upload Files", [
                "Open My Drive or the destination folder.",
                "Click New and choose File upload, or use the page upload control or drag files into the workspace.",
                "Select one or more files. Large files are divided into resumable chunks automatically.",
                "Watch the upload panel. You may continue browsing while server-side finalization completes.",
                "Refresh only if the interface asks you to; completed files appear in the destination folder.",
            ]),
            ("How To Upload A Folder", [
                "Open the destination folder.",
                "Click New and choose Folder upload.",
                "Select a local folder and confirm browser access when prompted.",
                "Keep the tab open while chunks transfer; final processing can continue after the transfer is accepted.",
                "Review the recreated folder structure after completion.",
            ]),
            ("How To Create And Organize Folders", [
                "Click New > New folder, enter a name, and create it.",
                "Open a folder by selecting its card or name.",
                "Use an item's actions menu to rename, move, or copy it.",
                "For drag-and-drop moves, drop the item on a destination folder or a folder target in the sidebar.",
                "Confirm the destination in My Drive after the operation finishes.",
            ]),
            ("How To Manage File Versions", [
                "Open the file actions and choose version history or details.",
                "Review archived versions and their dates.",
                "Upload a new version when the replacement should keep the same file record.",
                "Restore an earlier version when it should become current.",
                "Delete an archived version only when it is no longer required.",
            ]),
        ])

    add_section(doc, 2, "Recent",
        "Recent provides a short path back to files that the current user opened recently.",
        [
            "Recently opened files ordered by recent activity.",
            "The same preview and authorized item actions available from My Drive.",
            "Search, layout controls, and item details consistent with the main workspace.",
        ],
        [("How To Reopen Recent Work", [
            "Open Recent from the sidebar.",
            "Use search or visible metadata to locate the file.",
            "Open the file to launch its supported preview, or use Download when a local copy is needed.",
            "Star the file if it should remain easy to find after it drops out of the recent list.",
        ])])

    add_section(doc, 3, "Starred",
        "Starred is a personal favorites view for important files and folders.",
        [
            "Files and folders marked as favorites by the signed-in user.",
            "Normal item metadata, preview entry points, and permitted file actions.",
            "Grid or list display and search/filter controls shared with the drive workspace.",
        ],
        [("How To Add Or Remove A Star", [
            "Find the file or folder in My Drive, Recent, Shared, or another accessible view.",
            "Open its action menu and choose the star or favorite action.",
            "Open Starred to confirm the item appears.",
            "Run the same action again to remove the star without deleting or moving the item.",
        ])])

    add_section(doc, 4, "Shared",
        "Shared lists files and folders that another account explicitly shared with the signed-in user.",
        [
            "Shared files and folders, their owner, and the effective viewer or editor role.",
            "Direct navigation into shared folder contents.",
            "Preview and download actions when permitted.",
            "Editing actions only when the share grants editor access.",
        ],
        [
            ("How To Work With A Shared Item", [
                "Open Shared from the sidebar.",
                "Select a file to preview it or open a folder to browse its contents.",
                "Check the displayed permission before trying to change content.",
                "For viewer access, preview or download only as allowed.",
                "For editor access, use the available upload, create-folder, move, copy, rename, or version actions within the authorized scope.",
            ]),
            ("How To Share With A Registered User", [
                "Open the owner-controlled item actions and choose Share.",
                "Search for or select the recipient account.",
                "Choose viewer or editor access for the recipient.",
                "Save the sharing settings.",
                "Ask the recipient to open Shared after signing in.",
            ]),
        ])

    add_section(doc, 5, "Trash",
        "Trash holds deleted files and folders during the retention window so accidental deletions can be reversed.",
        [
            "Deleted item name, type, original owner or location context, and deletion date.",
            "Restore and permanent-delete actions for individual items.",
            "Bulk controls for selected trash items.",
            "Automatic purge behavior after the configured 30-day retention period.",
        ],
        [
            ("How To Restore An Item", [
                "Open Trash from the sidebar.",
                "Locate and select the file or folder.",
                "Choose Restore.",
                "Return to My Drive and confirm the restored item is available.",
            ]),
            ("How To Delete Permanently", [
                "Open Trash and select the item or items.",
                "Choose the permanent-delete or purge action.",
                "Review the confirmation carefully; this operation is not recoverable through the application.",
                "Confirm only when the content and any archived versions are no longer needed.",
            ]),
        ])

    add_section(doc, 6, "Storage",
        "Storage is the analytics and inventory view for understanding space use and finding large or recently changed files.",
        [
            "Total file count, folder count, and bytes used.",
            "File categories such as images, videos, audio, documents, spreadsheets, archives, code, and other.",
            "Search, category filtering, and sorting by modified date or size.",
            "File rows with preview and download entry points.",
            "System-wide totals for super administrators and account-scoped totals for standard users.",
        ],
        [("How To Review Storage Use", [
            "Open Storage from the sidebar footer.",
            "Review the headline totals and category breakdown.",
            "Choose a category when investigating a specific media type.",
            "Sort by Size to find the largest files or by Modified to review recent changes.",
            "Open a file preview before deciding whether to download, version, move, or delete it.",
        ])])

    add_section(doc, 7, "Users",
        "Users is the super-administrator workspace for controlling who can sign in and which account role they receive.",
        [
            "User name, email address, role, and creation details.",
            "Actions to add, edit, or delete an account.",
            "Role choices limited to Standard User and Super Admin.",
        ],
        [
            ("How To Add A User", [
                "Sign in as a super administrator.",
                "Open Users and click Add User.",
                "Enter the name, email, password, and role.",
                "Click Create User.",
                "Provide the sign-in details through an approved secure channel.",
            ]),
            ("How To Edit A User", [
                "Open Users and select Edit for the account.",
                "Update the name, email, or role.",
                "Enter a new password only when a reset is required; leave it blank to keep the current password.",
                "Save the account.",
            ]),
            ("How To Delete A User", [
                "Open Users and identify the correct account.",
                "Click Delete and review the confirmation.",
                "Do not delete your own currently signed-in account; the system blocks this action.",
                "Confirm only after deciding how the user's owned content should be handled operationally.",
            ]),
        ])

    add_section(doc, 8, "File Preview",
        "Preview opens supported multimedia and document formats inside the browser while preserving access controls and download options.",
        [
            "Images with gallery controls; PDFs; Word, Excel, and PowerPoint files; Markdown and source code.",
            "Video and audio players with range-based streaming.",
            "ZIP archive listings with entry preview or download.",
            "Design-file and unsupported-file fallback views.",
            "Back, download, and file-specific viewing controls.",
        ],
        [("How To Preview And Download A File", [
            "Open an accessible file from My Drive, Recent, Starred, Shared, Storage, or a valid sharing link.",
            "Use the viewer controls appropriate to the file type.",
            "For ZIP files, browse entries and open or download the required item.",
            "Click Download when the original file is needed locally and permission allows it.",
            "Use Back to return to the previous drive or sharing view.",
        ])])

    add_section(doc, 9, "Sharing Links",
        "Sharing links provide controlled external access to a file or folder without requiring every recipient to be a registered user.",
        [
            "General access set to Restricted or Anyone with the link.",
            "Viewer or Editor link role.",
            "Optional password, expiration date, and download permission.",
            "Recipient account or email selection for restricted access.",
            "Link copy, settings update, and revoke actions.",
        ],
        [
            ("How To Create Or Update A Sharing Link", [
                "Open the item's Share dialog.",
                "Choose Restricted for selected recipients or Anyone with the link for broad link access.",
                "Choose Viewer for read-only access or Editor for folder upload and folder-creation access.",
                "Set a password, expiration date, and download permission when required.",
                "Add registered recipients or recipient email addresses for restricted access.",
                "Save the settings and copy the generated link.",
            ]),
            ("How A Recipient Opens A Link", [
                "Open the received sharing URL.",
                "Sign in when the link is restricted to selected users.",
                "Enter the share password when prompted.",
                "Browse or preview the shared content.",
                "Download only when enabled; editor links may also upload files or create folders in a shared folder.",
            ]),
            ("How To Revoke A Link", [
                "Open the owner-controlled Share dialog for the item.",
                "Review the current link settings and recipients.",
                "Choose Revoke or delete the share.",
                "Confirm. The old token should no longer grant access.",
            ]),
        ])

    add_section(doc, 10, "Profile And Active Devices",
        "Profile manages the signed-in user's identity and password, while Active Devices manages database-backed sign-in sessions.",
        [
            "Profile name and account email.",
            "Current-password and new-password fields.",
            "Active sessions with device/browser information, IP address, and timing data.",
            "Controls to revoke one session or all other sessions.",
        ],
        [
            ("How To Update The Profile", [
                "Open the profile menu and choose Profile.",
                "Edit the display name.",
                "Save the profile and confirm the success message.",
            ]),
            ("How To Change The Password", [
                "Open Profile.",
                "Enter the current password.",
                "Enter and confirm a new password of at least eight characters.",
                "Save the change and use the new password at the next sign-in.",
            ]),
            ("How To Revoke A Device Session", [
                "Open the profile menu and choose Active Devices.",
                "Review the listed sessions and identify the session to remove.",
                "Revoke the selected session, or revoke all other sessions when account access is uncertain.",
                "Keep the current session only when the current device is trusted.",
            ]),
        ])

    add_text(doc, "System Controls And Good Practice", "Heading 1")
    add_text(doc, "Required Access And Validation", "Heading 2")
    add_bullets(doc, [
        "Every drive, preview, version, trash, and download operation checks ownership, super-admin status, or a valid share permission.",
        "Uploads are staged outside the public web directory and finalized server-side.",
        "Share passwords, expiration times, restricted recipients, link roles, and download settings are enforced before protected actions.",
        "Session cookies use account-backed expiration and can be revoked from Active Devices.",
        "Permanent deletion cannot be reversed through Trash; use it only after confirming retention requirements.",
    ])
    add_text(doc, "Recommended Operating Routine", "Heading 2")
    add_steps(doc, [
        "Start in My Drive and open the correct destination folder.",
        "Upload or create content, then wait for finalization to complete.",
        "Preview the file to confirm integrity and use versions for controlled replacements.",
        "Organize with folders, names, stars, move, and copy actions.",
        "Share with the least-privileged role, an expiry, and a password when sensitivity warrants it.",
        "Use Recent and Starred for daily retrieval, Storage for capacity checks, and Trash for recovery.",
        "Review Active Devices after password changes or whenever an unfamiliar session appears.",
    ])
    add_text(doc, "Complete Information Flow", "Heading 2")
    add_text(doc, "Multimedia system flow: Account sign-in -> My Drive access -> Upload/create -> Chunk transfer and server finalization -> File metadata and storage update -> Preview/version/organize -> Optional share to account or link -> Recent/Starred/Shared retrieval -> Storage review -> Trash retention or permanent purge.")

    doc.settings.update_fields_on_open = True
    doc.save(str(DOCX_OUT))
    return DOCX_OUT


PALETTE = {
    "blue": ("#E8EEFF", "#5B73E8"),
    "teal": ("#D9F8F2", "#159D8C"),
    "green": ("#DFF7E8", "#2FA66A"),
    "amber": ("#FFF2CC", "#E29A21"),
    "pink": ("#FFE1EB", "#E7508A"),
    "purple": ("#EFE2FF", "#9B5DE5"),
    "slate": ("#E8EDF4", "#718096"),
    "red": ("#FFE3E3", "#E85D5D"),
}


def add_arrow(d, x1, y1, x2, y2, color="#B8C6E8", width=2):
    d.add(Line(x1, y1, x2, y2, strokeColor=HexColor(color), strokeWidth=width))
    ang = math.atan2(y2 - y1, x2 - x1)
    s = 8
    pts = []
    for delta in (2.65, -2.65):
        pts.extend([x2 + s * math.cos(ang + delta), y2 + s * math.sin(ang + delta)])
    d.add(Polygon([x2, y2, *pts], fillColor=HexColor(color), strokeColor=HexColor(color)))


def wrap(text, max_chars):
    words = text.split()
    lines, cur = [], ""
    for word in words:
        candidate = word if not cur else cur + " " + word
        if len(candidate) > max_chars and cur:
            lines.append(cur)
            cur = word
        else:
            cur = candidate
    if cur:
        lines.append(cur)
    return lines


def centered_lines(d, text, cx, cy, max_chars=24, size=15, color="#344054", bold=False, leading=None):
    lines = wrap(text, max_chars)
    leading = leading or size * 1.25
    start = cy + (len(lines) - 1) * leading / 2
    for i, line in enumerate(lines):
        d.add(String(cx, start - i * leading, line, textAnchor="middle", fontName="Helvetica-Bold" if bold else "Helvetica", fontSize=size, fillColor=HexColor(color)))


def role_diagram(name, sign_label, modules, output_stem):
    width, height = 3400, 980
    d = Drawing(width, height)
    d.add(Rect(0, 0, width, height, fillColor=white, strokeColor=white))
    title_w, title_h = 300, 145
    tx = (width - title_w) / 2
    d.add(Rect(tx, 795, title_w, title_h, fillColor=HexColor("#10182B"), strokeColor=HexColor("#10182B")))
    centered_lines(d, f"{name} Role - How to Use Each Tab", width / 2, 867, 19, 25, "#DCE5FF", True, 31)
    d.add(Rect(width / 2 - 95, 725, 190, 48, rx=24, ry=24, fillColor=HexColor("#0F766E"), strokeColor=HexColor("#0F766E")))
    centered_lines(d, sign_label, width / 2, 749, 25, 15, "#FFFFFF", True)

    margin, gap = 55, 28
    module_w = (width - margin * 2 - gap * (len(modules) - 1)) / len(modules)
    top = 625
    for idx, (title, color_key, nodes) in enumerate(modules):
        x = margin + idx * (module_w + gap)
        fill, stroke = PALETTE[color_key]
        box_h = 350 if len(nodes) >= 4 else 300
        y = top - box_h
        add_arrow(d, width / 2, 725, x + module_w / 2, top, "#C9D5F3", 1.5)
        d.add(Rect(x, y, module_w, box_h, fillColor=HexColor(fill), strokeColor=HexColor(stroke), strokeWidth=2))
        centered_lines(d, title, x + module_w / 2, top - 25, 22, 16, stroke, True)
        node_h = 55
        usable_top = top - 62
        step = (box_h - 82) / max(1, len(nodes))
        for nidx, node in enumerate(nodes):
            ny = usable_top - (nidx + 1) * step
            node_fill, node_stroke = (PALETTE["red"] if node.startswith("Not available") or node.startswith("Cannot ") else ("#FFFFFF", stroke))
            d.add(Rect(x + 24, ny, module_w - 48, node_h, fillColor=HexColor(node_fill), strokeColor=HexColor(node_stroke), strokeWidth=1.5))
            centered_lines(d, node, x + module_w / 2, ny + node_h / 2, 30, 13, "#39465E", False, 15)
            if nidx > 0:
                add_arrow(d, x + module_w / 2, ny + node_h + max(2, step - node_h), x + module_w / 2, ny + node_h, stroke, 1.2)

    d.add(String(width - 35, 20, "Multimedia Management System", textAnchor="end", fontName="Helvetica", fontSize=12, fillColor=HexColor("#B3BAC6")))
    renderPDF.drawToFile(d, str(OUTPUT / f"{output_stem}.pdf"))
    renderSVG.drawToFile(d, str(OUTPUT / f"{output_stem}.svg"))
    return d


def overview_diagram():
    width, height = 2300, 1550
    d = Drawing(width, height)
    d.add(Rect(0, 0, width, height, fillColor=white, strokeColor=white))
    d.add(String(width / 2, height - 70, "Full System Overview Diagram", textAnchor="middle", fontName="Helvetica-Bold", fontSize=44, fillColor=HexColor("#111111")))
    d.add(Rect(85, height - 125, 145, 48, rx=24, ry=24, fillColor=HexColor("#0F766E"), strokeColor=HexColor("#0F766E")))
    centered_lines(d, "Sign in", 157, height - 101, 20, 16, "#FFFFFF", True)
    add_arrow(d, 157, height - 125, 157, height - 175, "#202020", 2)

    def box(x, y, w, h, title, subtitle, key="blue"):
        fill, stroke = PALETTE[key]
        d.add(Rect(x, y, w, h, fillColor=HexColor(fill), strokeColor=HexColor(stroke), strokeWidth=2))
        centered_lines(d, title, x + w / 2, y + h * 0.64, 27, 18, stroke, True)
        centered_lines(d, subtitle, x + w / 2, y + h * 0.32, 32, 14, "#465166", False, 16)
        return x + w / 2, y + h / 2

    my = box(55, 1165, 270, 120, "My Drive", "Upload, organize, preview, version, share", "blue")
    modules = [
        (70, 930, 250, 100, "New", "Folder, file upload, folder upload", "teal"),
        (370, 930, 250, 100, "Recent", "Recently opened files", "slate"),
        (670, 930, 250, 100, "Starred", "Favorite files and folders", "amber"),
        (970, 930, 250, 100, "Shared", "Account viewer/editor access", "green"),
        (1270, 930, 250, 100, "Trash", "Restore or purge retained items", "red"),
        (1570, 930, 250, 100, "Storage", "Usage, categories, large files", "teal"),
        (1870, 930, 250, 100, "Users", "Super-admin account control", "pink"),
    ]
    for spec in modules:
        c = box(*spec)
        add_arrow(d, my[0], 1165, c[0], 1030, "#333333", 1.5)

    preview = box(260, 650, 310, 115, "File Preview", "Image, PDF, Office, video, audio, code, ZIP", "purple")
    versions = box(670, 650, 270, 115, "Versions", "Upload, restore, or delete archived versions", "blue")
    share = box(1040, 650, 320, 115, "Sharing", "Registered recipients or configurable links", "green")
    queue = box(1460, 650, 300, 115, "Background Work", "Finalize uploads, ZIP folders, thumbnails", "teal")
    profile = box(1860, 650, 270, 115, "Profile / Devices", "Name, password, active sessions", "purple")
    for c in [preview, versions, share, queue, profile]:
        add_arrow(d, my[0], 1165, c[0], 765, "#333333", 1.5)

    viewer = box(850, 360, 260, 100, "Viewer Link", "Preview and optional download", "slate")
    editor = box(1160, 360, 260, 100, "Editor Link", "Upload and create folders", "green")
    restricted = box(1470, 360, 300, 100, "Access Controls", "Restricted/anyone, password, expiry", "red")
    add_arrow(d, share[0], 650, viewer[0], 460, "#333333", 1.5)
    add_arrow(d, share[0], 650, editor[0], 460, "#333333", 1.5)
    add_arrow(d, share[0], 650, restricted[0], 460, "#333333", 1.5)

    d.add(Rect(60, 95, 2050, 150, fillColor=HexColor("#F5F7FA"), strokeColor=HexColor("#94A3B8"), strokeWidth=2))
    d.add(String(90, 215, "Who can do what", fontName="Helvetica-Bold", fontSize=18, fillColor=HexColor("#475569")))
    legend = [
        ("Standard user", "Own items and shared items within granted permission", "slate"),
        ("Super administrator", "System-wide drive visibility plus user management", "pink"),
        ("Viewer recipient", "Read-only preview and allowed downloads", "blue"),
        ("Editor recipient", "Upload and create folders inside a shared folder", "green"),
    ]
    for i, (role, desc, key) in enumerate(legend):
        x = 90 + i * 500
        fill, stroke = PALETTE[key]
        d.add(Rect(x, 120, 450, 66, fillColor=HexColor(fill), strokeColor=HexColor(stroke)))
        d.add(String(x + 15, 158, role, fontName="Helvetica-Bold", fontSize=14, fillColor=HexColor(stroke)))
        centered_lines(d, desc, x + 280, 140, 33, 11, "#455166", False, 13)

    renderPDF.drawToFile(d, str(OUTPUT / "Full System Overview Diagram Multimedia Management System.pdf"))
    renderSVG.drawToFile(d, str(OUTPUT / "Full System Overview Diagram Multimedia Management System.svg"))


def build_diagrams():
    common = [
        ("My Drive", "blue", ["Browse folders and files", "Upload files or folders", "Preview, organize, and version", "Share or move to Trash"]),
        ("Recent", "slate", ["Review recently opened files", "Open preview or download"]),
        ("Starred", "amber", ["Review favorites", "Remove star when no longer needed"]),
        ("Shared", "green", ["Browse shared items", "Follow viewer/editor permission"]),
        ("Trash", "red", ["Restore retained items", "Permanently delete with care"]),
        ("Storage", "teal", ["Review usage totals", "Filter and sort files"]),
        ("Users", "pink", ["Manage accounts and roles"]),
        ("Profile", "purple", ["Update name or password", "Review active devices", "Sign out when finished"]),
    ]
    role_diagram("Super Admin", "Sign in as Super Admin", common, "Super Admin Role User Flow Diagram")

    user_modules = [tuple(m) for m in common]
    user_modules[6] = ("Users", "red", ["Not available to Standard User"])
    role_diagram("Standard User", "Sign in as Standard User", user_modules, "Standard User Role User Flow Diagram")

    viewer_modules = [
        ("Shared Link", "slate", ["Open valid sharing URL", "Sign in if restricted", "Enter password if required"]),
        ("Browse", "blue", ["View shared file or folder", "Open supported previews"]),
        ("Download", "amber", ["Download only when owner allows"]),
        ("Changes", "red", ["Cannot upload or create folders", "Cannot rename, move, or delete"]),
        ("Expiry", "purple", ["Access ends when link expires", "Owner may revoke at any time"]),
    ]
    role_diagram("Viewer Link", "Open Viewer Link", viewer_modules, "Viewer Link Role User Flow Diagram")

    editor_modules = [
        ("Shared Link", "green", ["Open valid sharing URL", "Sign in if restricted", "Enter password if required"]),
        ("Browse", "blue", ["View shared folder contents", "Open supported previews"]),
        ("Contribute", "teal", ["Upload files", "Create folders", "Track upload progress"]),
        ("Download", "amber", ["Download when owner allows"]),
        ("Limits", "red", ["Cannot manage owner accounts", "Cannot change link settings", "Cannot access unrelated content"]),
    ]
    role_diagram("Editor Link", "Open Editor Link", editor_modules, "Editor Link Role User Flow Diagram")
    overview_diagram()


if __name__ == "__main__":
    build_docx()
    build_diagrams()
    print(DOCX_OUT)
